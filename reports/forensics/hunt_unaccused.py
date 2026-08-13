#!/usr/bin/env python3
"""Download and scan Nature-family papers that were not in the Geng-named set.

Allegation generator only. Review hits by hand before writing them up.
"""
from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from download_source_data import ROOT, download_labeled
from scan_source_data import scan_paper

# Papers Geng / journals already named — do not treat as "unaccused".
ALREADY_NAMED = {
    "10.1038/s41586-024-08248-5",  # HDAC6
    "10.1038/s43018-023-00715-8",  # LGR4
    "10.1038/s41467-023-38842-6",  # Luo BAT
    "10.1038/s43587-024-00694-0",  # Lei Nature Aging
    "10.1038/s41586-026-10113-6",  # Qiu
    "10.1038/s41467-021-25561-z",  # Wang Qigang
    "10.1038/s41565-025-02082-0",  # Su
    "10.1038/s41556-026-01876-1",  # Kuang NCB
}

# Sibling / same-lab papers that Geng did not individually name, plus a
# small blind set from other mainland labs (WB/qPCR-heavy, Source Data).
PAPERS = [
    # Xiangya Luo/Zhou/Li corresponding — not the named BAT paper
    {
        "id": "grancalcin_nc",
        "doi": "10.1038/s41467-023-43787-x",
        "url": "https://www.nature.com/articles/s41467-023-43787-x",
        "lab": "xiangya_luo_zhou",
        "want": ["xlsx", "xls", "source"],
    },
    {
        "id": "temp_nc",
        "doi": "10.1038/s41467-023-39171-4",
        "url": "https://www.nature.com/articles/s41467-023-39171-4",
        "lab": "xiangya_zhou",
        "want": ["xlsx", "xls", "source"],
    },
    {
        "id": "dsba_nc",
        "doi": "10.1038/s41467-020-20665-4",
        "url": "https://www.nature.com/articles/s41467-020-20665-4",
        "lab": "xiangya_luo",
        "want": ["xlsx", "xls", "source"],
    },
    {
        "id": "circrna_nc",
        "doi": "10.1038/s41467-025-63343-z",
        "url": "https://www.nature.com/articles/s41467-025-63343-z",
        "lab": "xiangya_lei",
        "want": ["xlsx", "xls", "source"],
    },
    # Lei 2021 Nat Commun arthritis — older sibling, not Geng-named
    {
        "id": "lei_arthritis_nc",
        "doi": "10.1038/s41467-021-22454-z",
        "url": "https://www.nature.com/articles/s41467-021-22454-z",
        "lab": "xiangya_lei",
        "want": ["xlsx", "xls", "source"],
    },
    # Blind set: other mainland labs, metabolism/aging, Source Data present
    {
        "id": "fudan_fasn_nc",
        "doi": "10.1038/s41467-023-44393-7",
        "url": "https://www.nature.com/articles/s41467-023-44393-7",
        "lab": "fudan_zhao",
        "want": ["xlsx", "xls", "source"],
    },
    {
        "id": "sjtu_apc_nc",
        "doi": "10.1038/s41467-024-48914-w",
        "url": "https://www.nature.com/articles/s41467-024-48914-w",
        "lab": "sjtu_adipose",
        "want": ["xlsx", "xls", "source"],
    },
    {
        "id": "suda_cd36_nc",
        "doi": "10.1038/s41467-025-60437-6",
        "url": "https://www.nature.com/articles/s41467-025-60437-6",
        "lab": "suda_fudan",
        "want": ["xlsx", "xls", "source"],
        "skip": ["supplementary data", "supplementary table"],
    },
]


HIGH = {
    "exact_duplicate_rowgroup",
    "constant_offset",
    "shared_fractional_parts",
    "shared_fractional_parts_row",
    "affine_copy",
    "permutation_duplicate_vector",
    "within_group_duplicate",
    "precise_scalar_reuse",
    "exact_duplicate_vector",
}


def convert_xls(path: Path) -> Path | None:
    """xlrd 1.2 xls → xlsx so openpyxl can read it."""
    if path.suffix.lower() != ".xls":
        return path
    out = path.with_suffix(".xlsx")
    if out.exists() and out.stat().st_size > 1000:
        return out
    import xlrd
    from openpyxl import Workbook

    book = xlrd.open_workbook(str(path))
    wb = Workbook()
    wb.remove(wb.active)
    for si in range(book.nsheets):
        sh = book.sheet_by_index(si)
        ws = wb.create_sheet(title=(sh.name or f"sheet{si}")[:31])
        for r in range(min(sh.nrows, 400)):
            for c in range(min(sh.ncols, 80)):
                ws.cell(r + 1, c + 1, sh.cell_value(r, c))
    wb.save(out)
    print(f"  converted {path.name} -> {out.name}")
    return out


def scanable(folder: Path):
    files = []
    for p in folder.iterdir():
        if not p.is_file():
            continue
        if p.suffix.lower() not in {".xlsx", ".xls"}:
            continue
        # xls is converted with a 400×80 cap; allow larger binaries.
        if p.suffix.lower() == ".xlsx" and p.stat().st_size > 2_000_000:
            print(f"  skip large {p.name} {p.stat().st_size}")
            continue
        q = convert_xls(p)
        if q:
            files.append(q)
    return files


def main():
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    ROOT.mkdir(parents=True, exist_ok=True)
    summary = {}
    for spec in PAPERS:
        if spec["doi"] in ALREADY_NAMED:
            print("SKIP named", spec["doi"])
            continue
        kept = download_labeled(
            spec["id"],
            spec["url"],
            want_substr=spec.get("want"),
            skip_substr=spec.get("skip"),
        )
        folder = ROOT / spec["id"]
        paths = scanable(folder)
        print(f"  scanable n={len(paths)}")
        findings = scan_paper(spec["id"], paths)
        counts = Counter(f.get("type") for f in findings)
        high = [f for f in findings if f.get("type") in HIGH]
        summary[spec["id"]] = {
            "doi": spec["doi"],
            "lab": spec["lab"],
            "n_files": len(paths),
            "counts": dict(counts),
            "n_high": len(high),
            "high": high[:80],
        }
        print("  counts", dict(counts))
        for f in high[:25]:
            print(
                "   ",
                f.get("type"),
                f.get("file"),
                f.get("loc") or f.get("loc_a"),
                "||",
                f.get("loc_b"),
                "val",
                f.get("value"),
                "times",
                f.get("times"),
            )
            if f.get("vals_a"):
                print("      A", f["vals_a"][:10])
                print("      B", (f.get("vals_b") or [])[:10])
            if f.get("group"):
                print("      G", f["group"])
            la = (f.get("label") or f.get("label_a") or "")[:140]
            if la:
                print("      L", la)

    out = Path("/workspace/reports/forensics/unaccused_scan.json")
    out.write_text(json.dumps(summary, indent=2, default=str))
    print("WROTE", out)


if __name__ == "__main__":
    main()
